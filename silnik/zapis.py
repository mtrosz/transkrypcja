"""Zapis fragmentów transkrypcji do pliku tekstowego lub Worda, opcjonalnie ze znacznikami czasu."""
import os
from pathlib import Path

from docx import Document

from fragment import Fragment

PRZERWA_AKAPIT_S = 2.0
ROZSZERZENIA = {"txt": ".txt", "docx": ".docx"}


def akapity(fragmenty: list[Fragment]) -> list[str]:
    """Łączy fragmenty w akapity; nowy akapit po przerwie ≥ PRZERWA_AKAPIT_S."""
    grupy: list[list[str]] = []
    poprzedni_koniec = None
    for f in fragmenty:
        tekst = f.tekst.strip()
        if not tekst:
            continue
        if poprzedni_koniec is None or f.start - poprzedni_koniec >= PRZERWA_AKAPIT_S:
            grupy.append([])
        grupy[-1].append(tekst)
        poprzedni_koniec = f.koniec
    return [" ".join(g) for g in grupy]


def format_czasu(sekundy: float) -> str:
    godziny, reszta = divmod(int(sekundy), 3600)
    minuty, sek = divmod(reszta, 60)
    return f"{godziny:02d}:{minuty:02d}:{sek:02d}"


def wolna_sciezka(folder: Path, nazwa: str, rozszerzenie: str) -> Path:
    """Pierwsza nieistniejąca ścieżka: 'nazwa.ext', 'nazwa (2).ext', 'nazwa (3).ext'…"""
    kandydat = folder / f"{nazwa}{rozszerzenie}"
    numer = 2
    while kandydat.exists():
        kandydat = folder / f"{nazwa} ({numer}){rozszerzenie}"
        numer += 1
    return kandydat


def _zapisz_txt(sciezka: Path, fragmenty: list[Fragment], tytul: str) -> None:
    sciezka.write_text("\n\n".join(akapity(fragmenty)) + "\n", encoding="utf-8")


def _linie_z_czasem(fragmenty: list[Fragment]) -> list[str]:
    return [f"[{format_czasu(f.start)}] {f.tekst.strip()}" for f in fragmenty if f.tekst.strip()]


def _zapisz_txt_czas(sciezka: Path, fragmenty: list[Fragment], tytul: str) -> None:
    sciezka.write_text("\n".join(_linie_z_czasem(fragmenty)) + "\n", encoding="utf-8")


def _zapisz_docx_z(sciezka: Path, tytul: str, akapity_tekstu: list[str]) -> None:
    dokument = Document()
    dokument.add_heading(tytul, level=1)
    for akapit in akapity_tekstu:
        dokument.add_paragraph(akapit)
    dokument.save(str(sciezka))


def _zapisz_docx(sciezka: Path, fragmenty: list[Fragment], tytul: str) -> None:
    _zapisz_docx_z(sciezka, tytul, akapity(fragmenty))


def _zapisz_docx_czas(sciezka: Path, fragmenty: list[Fragment], tytul: str) -> None:
    _zapisz_docx_z(sciezka, tytul, _linie_z_czasem(fragmenty))


PISARZE = {"txt": _zapisz_txt, "docx": _zapisz_docx, "txt_czas": _zapisz_txt_czas, "docx_czas": _zapisz_docx_czas}


def _zapisz_w(folder: Path, fragmenty: list[Fragment], nagranie: Path, format: str, czas: bool) -> Path:
    cel = wolna_sciezka(folder, nagranie.stem, ROZSZERZENIA[format])
    tymczasowy = cel.with_name(f".{cel.name}.czesciowy")
    try:
        PISARZE[format + ("_czas" if czas else "")](tymczasowy, fragmenty, nagranie.stem)
        os.replace(tymczasowy, cel)
    except BaseException:
        tymczasowy.unlink(missing_ok=True)
        raise
    return cel


def zapisz(fragmenty: list[Fragment], nagranie: Path, format: str, biurko: Path, czas: bool = False) -> Path:
    """Zapisuje obok nagrania (albo na Biurku, gdy folder jest tylko do odczytu lub macOS odmawia zapisu). Nigdy nie nadpisuje."""
    if not os.access(nagranie.parent, os.W_OK):
        return _zapisz_w(biurko, fragmenty, nagranie, format, czas)
    try:
        return _zapisz_w(nagranie.parent, fragmenty, nagranie, format, czas)
    except PermissionError:
        return _zapisz_w(biurko, fragmenty, nagranie, format, czas)
