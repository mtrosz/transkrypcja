from pathlib import Path

import pytest
from docx import Document

import zapis
from fragment import Fragment
from zapis import akapity, format_czasu, wolna_sciezka, zapisz

FRAGMENTY = [
    Fragment(0, 2, "Dzień dobry."),
    Fragment(5, 7, "Art. 415 k.c."),
    Fragment(3725, 3727, "Koniec."),
]


def nagranie(folder: Path, nazwa: str = "Wykład 3.m4a") -> Path:
    p = folder / nazwa
    p.write_bytes(b"audio")
    return p


def test_akapity_laczy_bliskie_fragmenty():
    fr = [Fragment(0, 2, "Dzień dobry."), Fragment(2.5, 4, "Zaczynamy.")]
    assert akapity(fr) == ["Dzień dobry. Zaczynamy."]


def test_akapity_nowy_po_przerwie_2s():
    fr = [Fragment(0, 2, "Pierwszy."), Fragment(4.0, 5, "Drugi."), Fragment(5.5, 6, "Trzeci.")]
    assert akapity(fr) == ["Pierwszy.", "Drugi. Trzeci."]


def test_akapity_przerwa_tuz_ponizej_progu():
    assert akapity([Fragment(0, 2, "A."), Fragment(3.99, 5, "B.")]) == ["A. B."]


def test_akapity_przycina_spacje():
    fr = [Fragment(0, 1, " Ala "), Fragment(1, 2, "  ma kota.")]
    assert akapity(fr) == ["Ala ma kota."]


def test_akapity_puste():
    assert akapity([]) == []


@pytest.mark.parametrize(
    "sekundy, oczekiwane",
    [(0, "00:00:00"), (59.7, "00:00:59"), (252, "00:04:12"), (3723.9, "01:02:03"), (36000, "10:00:00")],
)
def test_format_czasu(sekundy, oczekiwane):
    assert format_czasu(sekundy) == oczekiwane


def test_wolna_sciezka_bez_kolizji(tmp_path):
    assert wolna_sciezka(tmp_path, "Wykład 3", ".txt") == tmp_path / "Wykład 3.txt"


def test_wolna_sciezka_kolejne_numery(tmp_path):
    (tmp_path / "Wykład 3.txt").write_text("x")
    (tmp_path / "Wykład 3 (2).txt").write_text("x")
    assert wolna_sciezka(tmp_path, "Wykład 3", ".txt") == tmp_path / "Wykład 3 (3).txt"


def test_zapisz_txt(tmp_path):
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path / "Biurko")
    assert wynik == tmp_path / "Wykład 3.txt"
    assert wynik.read_text(encoding="utf-8") == "Dzień dobry.\n\nArt. 415 k.c.\n\nKoniec.\n"


def test_zapisz_txt_czas(tmp_path):
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path / "Biurko", czas=True)
    assert wynik == tmp_path / "Wykład 3.txt"
    assert wynik.read_text(encoding="utf-8") == (
        "[00:00:00] Dzień dobry.\n[00:00:05] Art. 415 k.c.\n[01:02:05] Koniec.\n"
    )


def test_zapisz_docx_czas(tmp_path):
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "docx", biurko=tmp_path / "Biurko", czas=True)
    assert wynik == tmp_path / "Wykład 3.docx"
    doc = Document(str(wynik))
    assert [p.text for p in doc.paragraphs] == [
        "Wykład 3",
        "[00:00:00] Dzień dobry.",
        "[00:00:05] Art. 415 k.c.",
        "[01:02:05] Koniec.",
    ]
    assert doc.paragraphs[0].style.name == "Heading 1"


def test_zapisz_docx(tmp_path):
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "docx", biurko=tmp_path / "Biurko")
    assert wynik == tmp_path / "Wykład 3.docx"
    doc = Document(str(wynik))
    assert [p.text for p in doc.paragraphs] == ["Wykład 3", "Dzień dobry.", "Art. 415 k.c.", "Koniec."]
    assert doc.paragraphs[0].style.name == "Heading 1"


def test_zapisz_nie_nadpisuje(tmp_path):
    (tmp_path / "Wykład 3.txt").write_text("moje notatki", encoding="utf-8")
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path / "Biurko")
    assert wynik == tmp_path / "Wykład 3 (2).txt"
    assert (tmp_path / "Wykład 3.txt").read_text(encoding="utf-8") == "moje notatki"


def test_zapisz_na_biurko_gdy_folder_tylko_do_odczytu(tmp_path, monkeypatch):
    biurko = tmp_path / "Biurko"
    biurko.mkdir()
    folder = tmp_path / "nagrania"
    folder.mkdir()
    n = nagranie(folder)
    monkeypatch.setattr(zapis.os, "access", lambda sciezka, tryb: Path(sciezka) != folder)
    assert zapisz(FRAGMENTY, n, "txt", biurko=biurko) == biurko / "Wykład 3.txt"


def test_zapisz_sprzata_plik_czesciowy_przy_bledzie(tmp_path, monkeypatch):
    def psuj(sciezka, fragmenty, tytul):
        sciezka.write_text("pół pliku")
        raise RuntimeError("awaria")

    monkeypatch.setitem(zapis.PISARZE, "txt", psuj)
    with pytest.raises(RuntimeError):
        zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["Wykład 3.m4a"]


def test_zapisz_nazwa_z_kropkami_i_polskimi_znakami(tmp_path):
    n = nagranie(tmp_path, "Prawo cywilne – wykł. 3.1 „zobowiązania”.m4a")
    wynik = zapisz(FRAGMENTY, n, "txt", biurko=tmp_path)
    assert wynik.name == "Prawo cywilne – wykł. 3.1 „zobowiązania”.txt"


def test_zapisz_nagranie_bez_rozszerzenia(tmp_path):
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path, "Nagranie"), "txt", biurko=tmp_path)
    assert wynik == tmp_path / "Nagranie.txt"


def test_zapisz_sprzata_plik_czesciowy_przy_sigterm(tmp_path, monkeypatch):
    def przerwij(sciezka, fragmenty, tytul):
        sciezka.write_text("pół pliku")
        raise SystemExit(143)

    monkeypatch.setitem(zapis.PISARZE, "txt", przerwij)
    with pytest.raises(SystemExit):
        zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["Wykład 3.m4a"]


def test_zapisz_brak_uprawnien_w_folderze_nagrania_zapisuje_na_biurku(tmp_path, monkeypatch):
    biurko = tmp_path / "Biurko"
    biurko.mkdir()
    folder = tmp_path / "nagrania"
    folder.mkdir()
    n = nagranie(folder)
    oryginal = zapis.PISARZE["txt"]

    def pisarz(sciezka, fragmenty, tytul):
        oryginal(sciezka, fragmenty, tytul)
        if sciezka.parent == folder:
            raise PermissionError("Operation not permitted")

    monkeypatch.setitem(zapis.PISARZE, "txt", pisarz)
    assert zapisz(FRAGMENTY, n, "txt", biurko=biurko) == biurko / "Wykład 3.txt"
    assert sorted(p.name for p in folder.iterdir()) == ["Wykład 3.m4a"]
    assert sorted(p.name for p in biurko.iterdir()) == ["Wykład 3.txt"]


# --- wybrany folder zapisu ---


def test_zapisz_do_wybranego_folderu(tmp_path):
    cel = tmp_path / "Transkrypcje"
    cel.mkdir()
    (cel / "Wykład 3.txt").write_text("stary", encoding="utf-8")
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=tmp_path / "Biurko", folder=cel)
    assert wynik == cel / "Wykład 3 (2).txt"
    assert not (tmp_path / "Wykład 3.txt").exists()


def test_zapisz_nieistniejacy_folder_na_biurko(tmp_path):
    biurko = tmp_path / "Biurko"
    biurko.mkdir()
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "txt", biurko=biurko, folder=tmp_path / "odłączony pendrive")
    assert wynik == biurko / "Wykład 3.txt"


def test_zapisz_folder_tylko_do_odczytu_na_biurko(tmp_path, monkeypatch):
    biurko = tmp_path / "Biurko"
    biurko.mkdir()
    cel = tmp_path / "Tylko do odczytu"
    cel.mkdir()
    monkeypatch.setattr(zapis.os, "access", lambda sciezka, tryb: Path(sciezka) != cel)
    wynik = zapisz(FRAGMENTY, nagranie(tmp_path), "docx", biurko=biurko, folder=cel)
    assert wynik == biurko / "Wykład 3.docx"
